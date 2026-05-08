from pathlib import Path
from uuid import uuid4

from fastapi import HTTPException
from sqlalchemy.orm import Session
from sqlalchemy import or_, func
from docx import Document
from openpyxl import Workbook

from app.modules.document_libraries import repository
from app.modules.pages.models import Page
from app.modules.navigation.models import NavigationItem


BASE_DIR = Path(__file__).resolve().parent.parent.parent.parent
UPLOADS_DIR = BASE_DIR / "uploads"
DOCUMENTS_DIR = UPLOADS_DIR / "documents"

DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)


def create_library(db: Session, data):
    library = repository.create_library(
        db=db,
        title=data.title,
        description=data.description,
    )

    page = Page(
        portal_id=data.portal_id,
        title=data.title,
        description=data.description or "",
        status="published",
        is_home=False,
        is_visible=True,
        sort_order=0,
    )

    db.add(page)
    db.commit()
    db.refresh(page)

    navigation_item = NavigationItem(
        portal_id=data.portal_id,
        parent_id=data.parent_id,
        type="document_library",
        title=data.title,
        page_id=page.id,
        library_id=library.id,
        url=None,
        sort_order=0,
        is_visible=True,
        icon=None,
        icon_type=None,
        icon_file_url=None,
        color=None,
        is_bold=False,
        is_italic=False,
    )

    db.add(navigation_item)
    db.commit()
    db.refresh(navigation_item)

    return library


def get_libraries(db: Session):
    return repository.get_libraries(db)


def create_folder(db: Session, library_id: int, data):
    library = repository.get_library_by_id(db, library_id)

    if not library:
        raise HTTPException(status_code=404, detail="Библиотека не найдена")

    title = data.title.strip()

    if not title:
        raise HTTPException(status_code=400, detail="Название папки обязательно")

    if data.parent_id:
        parent = repository.get_document_by_id(db, data.parent_id)

        if not parent or not parent.is_folder:
            raise HTTPException(status_code=404, detail="Родительская папка не найдена")

        if parent.library_id != library_id:
            raise HTTPException(status_code=400, detail="Папка относится к другой библиотеке")

    return repository.create_folder(
        db=db,
        library_id=library_id,
        title=title,
        parent_id=data.parent_id,
    )


def create_document(db: Session, library_id: int, data):
    library = repository.get_library_by_id(db, library_id)

    if not library:
        raise HTTPException(status_code=404, detail="Библиотека не найдена")

    if data.parent_id:
        parent = repository.get_document_by_id(db, data.parent_id)

        if not parent or not parent.is_folder:
            raise HTTPException(status_code=404, detail="Родительская папка не найдена")

        if parent.library_id != library_id:
            raise HTTPException(status_code=400, detail="Папка относится к другой библиотеке")

    document_type = data.document_type.lower().strip()

    if document_type not in ["word", "excel"]:
        raise HTTPException(
            status_code=400,
            detail="Допустимые типы документов: word, excel",
        )

    safe_title = data.title.strip() or "Новый документ"

    if document_type == "word":
        filename = f"{uuid4()}.docx"
        file_path = DOCUMENTS_DIR / filename

        doc = Document()
        doc.add_heading(safe_title, level=1)
        doc.save(file_path)

    elif document_type == "excel":
        filename = f"{uuid4()}.xlsx"
        file_path = DOCUMENTS_DIR / filename

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Лист 1"
        sheet["A1"] = safe_title
        workbook.save(file_path)

    relative_path = f"/uploads/documents/{filename}"

    return repository.create_document(
        db=db,
        library_id=library_id,
        parent_id=data.parent_id,
        title=safe_title,
        document_type=document_type,
        file_path=relative_path,
        original_filename=filename,
        is_folder=False,
        created_by="Михаил",
    )


def upload_document(db: Session, library_id: int, file, parent_id: int | None = None):
    library = repository.get_library_by_id(db, library_id)

    if not library:
        raise HTTPException(status_code=404, detail="Библиотека не найдена")

    if parent_id:
        parent = repository.get_document_by_id(db, parent_id)

        if not parent or not parent.is_folder:
            raise HTTPException(status_code=404, detail="Родительская папка не найдена")

        if parent.library_id != library_id:
            raise HTTPException(status_code=400, detail="Папка относится к другой библиотеке")

    original_filename = file.filename or "uploaded_file"
    saved_filename = f"{uuid4()}_{original_filename}"
    file_path = DOCUMENTS_DIR / saved_filename

    with open(file_path, "wb") as buffer:
        buffer.write(file.file.read())

    relative_path = f"/uploads/documents/{saved_filename}"

    extension = Path(original_filename).suffix.lower().replace(".", "")
    document_type = extension or "file"

    return repository.create_document(
        db=db,
        library_id=library_id,
        parent_id=parent_id,
        title=original_filename,
        document_type=document_type,
        file_path=relative_path,
        original_filename=original_filename,
        is_folder=False,
        created_by="Михаил",
    )


def delete_file_from_disk(document):
    if not document.file_path:
        return

    try:
        full_path = BASE_DIR / document.file_path.lstrip("/")
        if full_path.exists():
            full_path.unlink()
    except Exception:
        pass


def delete_document_recursive(db: Session, document):
    children = repository.get_children(db, document.id)

    for child in children:
        delete_document_recursive(db, child)

    if not document.is_folder:
        delete_file_from_disk(document)

    db.delete(document)


def delete_document(
    db: Session,
    document_id: int,
    mode: str = "folder_only",
):
    document = repository.get_document_by_id(db, document_id)

    if not document:
        raise HTTPException(status_code=404, detail="Документ не найден")

    if mode not in ["folder_only", "with_children"]:
        raise HTTPException(status_code=400, detail="Неверный режим удаления")

    if document.is_folder:
        children = repository.get_children(db, document.id)

        if children and mode == "folder_only":
            for child in children:
                child.parent_id = document.parent_id

            db.delete(document)
            db.commit()

            return {"message": "Папка удалена, содержимое перемещено на уровень выше"}

        if children and mode == "with_children":
            delete_document_recursive(db, document)
            db.commit()

            return {"message": "Папка удалена вместе с содержимым"}

        db.delete(document)
        db.commit()

        return {"message": "Папка удалена"}

    delete_file_from_disk(document)

    db.delete(document)
    db.commit()

    return {"message": "Документ удалён"}


def rename_document(db: Session, document_id: int, title: str):
    document = repository.update_document_title(db, document_id, title.strip())

    if not document:
        raise HTTPException(status_code=404, detail="Документ не найден")

    return document


def get_documents_by_library(
    db: Session,
    library_id: int,
    parent_id: int | None = None,
    limit: int = 50,
    offset: int = 0,
):
    library = repository.get_library_by_id(db, library_id)

    if not library:
        raise HTTPException(status_code=404, detail="Библиотека не найдена")

    query = repository.get_documents_query(db, library_id, parent_id)

    total = query.count()

    items = (
        query
        .order_by(
            repository.LibraryDocument.is_folder.desc(),
            repository.LibraryDocument.id.desc(),
        )
        .offset(offset)
        .limit(limit)
        .all()
    )

    return {
        "items": items,
        "total": total,
        "limit": limit,
        "offset": offset,
    }


def search_documents(
    db: Session,
    library_id: int,
    query: str,
    limit: int = 200,
    offset: int = 0,
):
    library = repository.get_library_by_id(db, library_id)

    if not library:
        raise HTTPException(status_code=404, detail="Библиотека не найдена")

    search_text = (query or "").strip()

    if not search_text:
        return get_documents_by_library(
            db=db,
            library_id=library_id,
            parent_id=None,
            limit=limit,
            offset=offset,
        )

    pattern = f"%{search_text}%"

    search_query = (
        db.query(repository.LibraryDocument)
        .filter(repository.LibraryDocument.library_id == library_id)
        .filter(
            or_(
                repository.LibraryDocument.title.ilike(pattern),
                repository.LibraryDocument.document_type.ilike(pattern),
                repository.LibraryDocument.original_filename.ilike(pattern),
                repository.LibraryDocument.created_by.ilike(pattern),
                func.to_char(
                    repository.LibraryDocument.created_at,
                    "DD.MM.YYYY HH24:MI"
                ).ilike(pattern),
                func.to_char(
                    repository.LibraryDocument.updated_at,
                    "DD.MM.YYYY HH24:MI"
                ).ilike(pattern),
            )
        )
    )

    total = search_query.count()

    items = (
        search_query
        .order_by(
            repository.LibraryDocument.is_folder.desc(),
            repository.LibraryDocument.updated_at.desc(),
            repository.LibraryDocument.id.desc(),
        )
        .offset(offset)
        .limit(limit)
        .all()
    )

    return {
        "items": items,
        "total": total,
        "limit": limit,
        "offset": offset,
    }


def move_document(
    db: Session,
    document_id: int,
    parent_id: int | None,
):
    document = repository.get_document_by_id(db, document_id)

    if not document:
        raise HTTPException(status_code=404, detail="Документ не найден")

    if parent_id is None:
        document.parent_id = None
        db.commit()
        db.refresh(document)
        return document

    target = repository.get_document_by_id(db, parent_id)

    if not target or not target.is_folder:
        raise HTTPException(status_code=400, detail="Целевая папка не найдена")

    if target.library_id != document.library_id:
        raise HTTPException(status_code=400, detail="Разные библиотеки")

    if document.id == target.id:
        raise HTTPException(status_code=400, detail="Нельзя переместить в самого себя")

    def is_descendant(child, parent_id):
        while child.parent_id:
            if child.parent_id == parent_id:
                return True
            child = repository.get_document_by_id(db, child.parent_id)
            if not child:
                break
        return False

    if document.is_folder:
        if is_descendant(target, document.id):
            raise HTTPException(
                status_code=400,
                detail="Нельзя переместить папку внутрь своей вложенности",
            )

    document.parent_id = parent_id

    db.commit()
    db.refresh(document)

    return document